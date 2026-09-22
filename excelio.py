"""
Excel output — one workbook, three sheets:

  "LinkedIn"              — LinkedIn jobs (incl. hiring-team link column)
  "Company Career Pages"  — ATS + Workday jobs, collated across companies
  "All Sources"           — remote aggregators + India boards

Links are written as real, clickable hyperlinks. Tracking columns
(Applied?, QC status, Notes) start blank for your pipeline management.
The same column helpers are reused by the master tracker (tracker.py).
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="0563C1", underline="single")

# (Header, job-dict key, column width). Keys prefixed "_" are blank/tracking
# or filled by the writer (date). "salary_lpa" is hidden helper, not shown.
LINKEDIN_COLUMNS = [
    ("Company name", "company", 26),
    ("Job role name", "title", 34),
    ("Job link", "job_link", 44),
    ("Job ID", "job_id", 14),
    ("Hiring-team profile link", "hiring_team_link", 26),
    ("Source", "source", 12),
    ("Role type", "role_type", 14),
    ("Location", "location", 24),
    ("Salary", "salary", 18),
    ("Remote", "remote", 9),
    ("CV match score", "match_score", 13),
    ("Eligibility", "eligibility", 12),
    ("Fit reason", "fit_reason", 40),
    ("Fit score", "fit_score", 10),
    ("First seen", "_fetched", 18),
    ("Applied?", "_applied", 10),
    ("QC status", "_qc", 12),
    ("Notes", "_notes", 30),
]

# Career + All-sources share the same shape (no hiring-team column).
_COMMON_COLUMNS = [
    ("Company name", "company", 26),
    ("Job role name", "title", 34),
    ("Job link", "job_link", 44),
    ("Job ID", "job_id", 20),
    ("Source", "source", 18),
    ("Role type", "role_type", 14),
    ("Location", "location", 24),
    ("Salary", "salary", 18),
    ("Remote", "remote", 9),
    ("CV match score", "match_score", 13),
    ("Eligibility", "eligibility", 12),
    ("Fit reason", "fit_reason", 40),
    ("Fit score", "fit_score", 10),
    ("First seen", "_fetched", 18),
    ("Applied?", "_applied", 10),
    ("QC status", "_qc", 12),
    ("Notes", "_notes", 30),
]
CAREER_COLUMNS = _COMMON_COLUMNS
ALL_SOURCES_COLUMNS = _COMMON_COLUMNS

_LINK_KEYS = {"job_link", "hiring_team_link"}
_BLANK_KEYS = {"_applied", "_qc", "_notes"}


def write_sheet(ws, columns, jobs, fetched_at):
    """Write a formatted sheet. Reused by the master tracker too."""
    for col_idx, (header, _key, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    jobs = sorted(jobs, key=lambda j: j.get("match_score", 0), reverse=True)

    for row_idx, job in enumerate(jobs, start=2):
        for col_idx, (_header, key, _w) in enumerate(columns, 1):
            if key == "_fetched":
                value = job.get("_fetched") or fetched_at
            elif key in _BLANK_KEYS:
                value = job.get(key, "")   # preserve if merged from master
            elif key == "remote":
                value = "Yes" if job.get("remote") else "No"
            else:
                value = job.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if key in _LINK_KEYS and value:
                cell.hyperlink = value
                cell.font = _LINK_FONT

    if jobs:
        last_col = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{last_col}{len(jobs) + 1}"


def write_workbook(linkedin_jobs, career_jobs, other_jobs, fetched_at,
                   out_dir=None):
    """Write the per-run workbook (three sheets) and return its path."""
    out_dir = out_dir or config.OUTPUT_DIR
    os.makedirs(out_dir, exist_ok=True)

    wb = Workbook()
    ws_li = wb.active
    ws_li.title = "LinkedIn"
    write_sheet(ws_li, LINKEDIN_COLUMNS, linkedin_jobs, fetched_at)

    ws_cp = wb.create_sheet("Company Career Pages")
    write_sheet(ws_cp, CAREER_COLUMNS, career_jobs, fetched_at)

    ws_all = wb.create_sheet("All Sources")
    write_sheet(ws_all, ALL_SOURCES_COLUMNS, other_jobs, fetched_at)

    safe_stamp = fetched_at.replace(":", "").replace(" ", "_").replace("-", "")
    path = os.path.join(out_dir, f"jobs_{safe_stamp}.xlsx")
    wb.save(path)
    return path


def write_whats_new(new_jobs, fetched_at, date_stamp, out_dir=None):
    """
    Write a one-sheet "what's new today" workbook of just the jobs added to the
    master this run. `new_jobs` are the dicts returned by tracker.merge (each
    tagged with "_section"). Returns the path, or "" if there were no new jobs.

    Used by the headless/scheduled mode so a daily run leaves a compact digest.
    """
    if not new_jobs:
        return ""
    out_dir = out_dir or config.OUTPUT_DIR
    os.makedirs(out_dir, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "New today"
    # Reuse the common column shape (no hiring-team column, matches career/all).
    write_sheet(ws, _COMMON_COLUMNS, new_jobs, fetched_at)

    path = os.path.join(out_dir, f"whats_new_{date_stamp}.xlsx")
    wb.save(path)
    return path
