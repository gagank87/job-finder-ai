"""
Master tracker — a single, ever-growing, de-duplicated Excel file.

Every run merges its newly found jobs into `master_tracker.xlsx`:
  * A job already present (same source+id, or same normalized link) is NOT
    added again — no duplicates, ever.
  * Genuinely new jobs are appended with a "First seen" date.
  * Your manual edits (Applied?, QC status, Notes) are PRESERVED across runs
    — the merge never overwrites existing rows.

This is what makes repeated / scheduled runs safe: the master accumulates
the universe of jobs you've seen, and each run only surfaces what's new.
"""

import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

_HEADER_FILL = PatternFill("solid", fgColor="215968")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="0563C1", underline="single")

# The master's single-sheet schema. "Section" marks where each job came from.
MASTER_COLUMNS = [
    ("Section", 16),
    ("Company name", 26),
    ("Job role name", 34),
    ("Job link", 46),
    ("Job ID", 18),
    ("Source", 18),
    ("Role type", 14),
    ("Location", 24),
    ("Salary", 18),
    ("Remote", 9),
    ("CV match score", 13),
    ("Eligibility", 12),
    ("Fit reason", 40),
    ("Fit score", 10),
    ("Hiring-team profile link", 24),
    ("First seen", 18),
    ("Applied?", 10),
    ("QC status", 12),
    ("Notes", 30),
]
_LINK_COL_HEADERS = {"Job link", "Hiring-team profile link"}


def _norm_link(link):
    return (link or "").split("?")[0].rstrip("/").lower()


def _job_keys(source, job_id, link):
    """Identity keys for dedupe: primary source+id, secondary normalized link."""
    keys = set()
    if job_id:
        keys.add(f"id::{str(source).split(' ')[0].lower()}::{job_id}")
    nl = _norm_link(link)
    if nl:
        keys.add(f"link::{nl}")
    return keys


def _row_from_job(job, section, fetched_at):
    return {
        "Section": section,
        "Company name": job.get("company", ""),
        "Job role name": job.get("title", ""),
        "Job link": job.get("job_link", ""),
        "Job ID": job.get("job_id", ""),
        "Source": job.get("source", ""),
        "Role type": job.get("role_type", ""),
        "Location": job.get("location", ""),
        "Salary": job.get("salary", ""),
        "Remote": "Yes" if job.get("remote") else "No",
        "CV match score": job.get("match_score", 0),
        "Eligibility": job.get("eligibility", ""),
        "Fit reason": job.get("fit_reason", ""),
        "Fit score": job.get("fit_score", ""),
        "Hiring-team profile link": job.get("hiring_team_link", ""),
        "First seen": fetched_at,
        "Applied?": "",
        "QC status": "",
        "Notes": "",
    }


def _load_existing(path):
    """Return (list_of_row_dicts, set_of_keys) from an existing master."""
    rows, keys = [], set()
    if not os.path.exists(path):
        return rows, keys
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, r))
        rows.append(row)
        keys |= _job_keys(row.get("Source", ""), row.get("Job ID", ""),
                          row.get("Job link", ""))
    return rows, keys


def merge(all_jobs_by_section, fetched_at, path=None):
    """
    Merge new jobs into the master tracker.

    all_jobs_by_section: dict like {"LinkedIn": [...], "Career Pages": [...],
                                    "All Sources": [...]}.
    Returns (added_count, total_in_master, new_jobs) where new_jobs is the list
    of the genuinely-new job dicts added this run (each tagged with its
    "Section") — used by the headless "what's new today" digest.
    """
    path = path or config.MASTER_TRACKER
    existing_rows, existing_keys = _load_existing(path)

    added = 0
    new_jobs = []
    seen_this_run = set(existing_keys)
    for section, jobs in all_jobs_by_section.items():
        for job in jobs:
            keys = _job_keys(job.get("source", ""), job.get("job_id", ""),
                             job.get("job_link", ""))
            if keys & seen_this_run:
                continue  # already in master OR already added this run
            existing_rows.append(_row_from_job(job, section, fetched_at))
            new_jobs.append({**job, "_section": section})
            seen_this_run |= keys
            added += 1

    # Sort: newest-fit first is less important here; keep by match score desc.
    existing_rows.sort(key=lambda r: (r.get("CV match score") or 0), reverse=True)
    _write(path, existing_rows)
    return added, len(existing_rows), new_jobs


def _append_note(existing, addition):
    """Append a note to a cell's existing text, separated by ' | '."""
    existing = str(existing or "").strip()
    addition = str(addition or "").strip()
    if not addition:
        return existing
    if not existing:
        return addition
    return f"{existing} | {addition}"


def _update_matching_rows(jobs, updates, note, path):
    """
    Match tracker rows to `jobs` by the same identity keys used for dedupe and
    apply `updates` (a dict of column -> value) plus append `note` to Notes.

    Returns (matched_count, path). Rows not matching any job are left untouched,
    preserving all your manual edits. Raises PermissionError if the workbook is
    open/locked (caller handles it like elsewhere).
    """
    path = path or config.MASTER_TRACKER
    rows, _keys = _load_existing(path)
    if not rows:
        return 0, path

    # Build the set of identity keys we're looking to update.
    target_keys = set()
    for job in jobs:
        target_keys |= _job_keys(job.get("source", ""), job.get("job_id", ""),
                                 job.get("job_link", ""))
    if not target_keys:
        return 0, path

    matched = 0
    for row in rows:
        row_keys = _job_keys(row.get("Source", ""), row.get("Job ID", ""),
                             row.get("Job link", ""))
        if not (row_keys & target_keys):
            continue
        for col, val in updates.items():
            row[col] = val
        if note:
            row["Notes"] = _append_note(row.get("Notes"), note)
        matched += 1

    if matched:
        _write(path, rows)
    return matched, path


def mark_prepared(jobs, note="", path=None):
    """
    Mark the given jobs as 'Prepared' in the master tracker.

    Sets QC status = "Prepared" and appends `note` (e.g. the output folder) to
    Notes. Does NOT touch Applied? — a prepared job hasn't been submitted.
    Returns (matched_count, path). Raises PermissionError if the file is locked.
    """
    return _update_matching_rows(
        jobs, {"QC status": "Prepared"}, note, path)


def mark_applied(jobs, note="", path=None):
    """
    Mark the given jobs as truly 'Applied' in the master tracker.

    Called ONLY after autosubmit.submit returned a verified confirmation. Sets
    Applied? = "Yes", QC status = "Auto-applied", and appends the confirmation
    reference to Notes. Returns (matched_count, path). Raises PermissionError if
    the file is locked.
    """
    return _update_matching_rows(
        jobs, {"Applied?": "Yes", "QC status": "Auto-applied"}, note, path)


def _write(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Tracker"
    for col_idx, (header, width) in enumerate(MASTER_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (header, _w) in enumerate(MASTER_COLUMNS, 1):
            value = row.get(header, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if header in _LINK_COL_HEADERS and value:
                cell.hyperlink = value
                cell.font = _LINK_FONT

    if rows:
        last_col = get_column_letter(len(MASTER_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"
    wb.save(path)
